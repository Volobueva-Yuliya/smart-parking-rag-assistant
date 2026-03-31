from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
import re

def create_presentation(md_file_path, output_pptx_path):
    prs = Presentation()

    with open(md_file_path, 'r', encoding='utf-8') as f:
        content = f.read()

    # Split by "### Slide X:"
    slides_content = re.split(r'### Slide \d+:', content)
    # The first part is the header, skip it
    slides_content = [s.strip() for s in slides_content[1:]]

    # Slide 1: Title Slide
    if len(slides_content) >= 1:
        title_content = slides_content[0]
        slide_layout = prs.slide_layouts[0] # Title Slide
        slide = prs.slides.add_slide(slide_layout)
        
        # Extract title and subtitle
        lines = title_content.split('\n')
        title_text = ""
        subtitle_text = []
        for line in lines:
            line = line.strip('- ').strip('*')
            if "Project Name:" in line:
                title_text = line.replace("Project Name:", "").strip()
            elif "Subtitle:" in line:
                subtitle_text.append(line.replace("Subtitle:", "").strip())
            else:
                subtitle_text.append(line)
        
        slide.shapes.title.text = title_text
        slide.placeholders[1].text = "\n".join(subtitle_text)

    # Content Slides (2 to 12)
    for i in range(1, len(slides_content)):
        content_text = slides_content[i]
        slide_layout = prs.slide_layouts[1] # Title and Content
        slide = prs.slides.add_slide(slide_layout)
        
        lines = content_text.split('\n')
        title_line = lines[0].strip('**').strip()
        body_lines = lines[1:]
        
        slide.shapes.title.text = title_line
        
        tf = slide.placeholders[1].text_frame
        tf.word_wrap = True
        
        for line in body_lines:
            line = line.strip()
            if not line:
                continue
            
            # Handle bullet points
            p = tf.add_paragraph()
            
            # Determine indentation level
            indent_match = re.match(r'^\s*(\d+\.|-|\*)\s+', line)
            indent_level = 0
            if line.startswith('  ') or line.startswith('\t'):
                indent_level = 1
            
            clean_line = re.sub(r'^\s*(\d+\.|-|\*)\s+', '', line).strip()
            clean_line = clean_line.replace('**', '')
            
            p.text = clean_line
            p.level = indent_level

            # Specific logic for Evaluation Results (Slide 9)
            if "Evaluation Results" in title_line:
                if any(metric in clean_line for metric in ["Recall@3:", "Recall@1:", "Latency:"]):
                    p.font.bold = True

    prs.save(output_pptx_path)
    print(f"Presentation saved to {output_pptx_path}")

if __name__ == "__main__":
    create_presentation("presentation_slides.md", "Smart_Parking_Assistant_Presentation.pptx")
